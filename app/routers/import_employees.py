"""
Import Excel des employés (fichier export eRH).
Format attendu : onglet « Employés » avec les colonnes du fichier export.
"""
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session
from datetime import datetime, date
import io

import openpyxl

from ..database import get_db
from ..models.employee_local import EmployeeLocal
from ..models.user import User
from ..services.auth_service import require_editor

router = APIRouter(prefix="/api/employees", tags=["Employés (import local)"])

# Correspondance header Excel (en minuscule) → champ du modèle
_HEADER_MAP = {
    "matricule":          "matricule",
    "nom":                "nom",
    "prénom":             "prenom",
    "prenom":             "prenom",
    "email":              "email",
    "fonction":           "fonction",
    "service":            "service",
    "type contrat":       "type_contrat",
    "statut":             "statut",
    "date embauche":      "date_embauche",
    "date fin cdd":       "date_fin_cdd",
    "fin période essai":  "fin_periode_essai",
    "fin periode essai":  "fin_periode_essai",
    "genre":              "genre",
    "téléphone":          "telephone",
    "telephone":          "telephone",
    "manager niveau 1":   "manager_n1",
    "manager niveau 2":   "manager_n2",
}

_DATE_FIELDS = {"date_embauche", "date_fin_cdd", "fin_periode_essai"}


def _to_str(val) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    return None if s in ("", "nan", "None", "—", "-", "NaT") else s


def _to_date(val) -> date | None:
    if val is None:
        return None
    if isinstance(val, date) and not isinstance(val, datetime):
        return val
    if isinstance(val, datetime):
        return val.date()
    s = str(val).strip()
    if not s or s in ("nan", "None", "—", "-", "NaT"):
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    return None


def _format_matricule(raw) -> str | None:
    if raw is None:
        return None
    s = str(raw).strip()
    if not s or s in ("nan", "None"):
        return None
    try:
        return str(int(float(s))).zfill(4)
    except (ValueError, TypeError):
        return s if s else None


@router.post("/import", status_code=200, summary="Importer les employés depuis le fichier Excel export")
async def import_employees(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: User = Depends(require_editor),
):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(400, "Seuls les fichiers .xlsx sont acceptés")

    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)

    # Trouver l'onglet (flexible sur le nom exact)
    sheet_name = next(
        (n for n in wb.sheetnames if "employ" in n.lower()),
        None,
    )
    if sheet_name is None:
        raise HTTPException(
            400,
            f"Onglet 'Employés' introuvable. Onglets présents : {wb.sheetnames}",
        )

    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, "Feuille vide")

    # Construire l'index des colonnes à partir de la ligne d'en-tête
    raw_headers = [str(h or "").strip().lower() for h in rows[0]]
    col_idx: dict[str, int] = {}
    for i, h in enumerate(raw_headers):
        field = _HEADER_MAP.get(h)
        if field and field not in col_idx:
            col_idx[field] = i

    missing = {"matricule", "nom"} - col_idx.keys()
    if missing:
        raise HTTPException(400, f"Colonnes obligatoires manquantes : {missing}")

    created = updated = 0
    errors: list[str] = []

    for row_num, row in enumerate(rows[1:], start=2):
        matricule = _format_matricule(row[col_idx["matricule"]] if "matricule" in col_idx else None)
        if not matricule:
            continue  # Ligne vide

        nom = _to_str(row[col_idx["nom"]] if "nom" in col_idx else None)
        if not nom:
            errors.append(f"Ligne {row_num} : nom manquant (matricule {matricule})")
            continue

        data: dict = {"matricule": matricule, "nom": nom}
        for field, idx in col_idx.items():
            if field in ("matricule", "nom"):
                continue
            raw = row[idx] if idx < len(row) else None
            if field in _DATE_FIELDS:
                data[field] = _to_date(raw)
            else:
                data[field] = _to_str(raw)

        existing = (
            db.query(EmployeeLocal)
            .filter(EmployeeLocal.matricule == matricule)
            .first()
        )
        if existing:
            for k, v in data.items():
                setattr(existing, k, v)
            updated += 1
        else:
            db.add(EmployeeLocal(**data))
            created += 1

    db.commit()
    return {
        "message": "Import terminé",
        "created": created,
        "updated": updated,
        "errors_count": len(errors),
        "errors": errors[:20],
    }


@router.get("/local", summary="Liste des employés importés localement")
def list_employees_local(
    db: Session = Depends(get_db),
    _: User = Depends(require_editor),
):
    return db.query(EmployeeLocal).order_by(EmployeeLocal.nom).all()
