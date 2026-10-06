"""
Import Excel des licences (fichier multi-onglets).
Onglets traités :
  - Licences      → licences_kaspersky
  - Machines      → kaspersky_machines
  - M365 Comptes  → licences_m365_comptes
  - M365 Membres  → licences_m365_membres
  - Adobe         → licences_adobe
  - Power BI      → licences_powerbi
"""
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session
from datetime import datetime, date
import io, math

import openpyxl

from ..database import get_db
from ..models.licence_fichier import (
    LicenceKaspersky, KasperskyMachine,
    LicenceM365Compte, LicenceM365Membre,
    LicenceAdobe, LicencePowerBI,
)
from ..models.user import User
from ..services.auth_service import require_editor

router = APIRouter(prefix="/api/licences", tags=["Licences (import fichier)"])


# ── Utilitaires ───────────────────────────────────────────────────────────────

def _to_str(val) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    return None if s in ("", "nan", "None", "—", "-", "NaT") else s


def _to_date(val) -> date | None:
    if val is None:
        return None
    if isinstance(val, date) and not isinstance(val, datetime):
        return val
    if isinstance(val, datetime):
        return val.date()
    s = str(val).strip()
    if not s or s in ("nan", "None", "—", "-", "NaT"):
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    return None


def _to_int(val) -> int | None:
    if val is None:
        return None
    try:
        f = float(str(val))
        return None if math.isnan(f) else int(f)
    except (ValueError, TypeError):
        return None


def _is_valid_num(val) -> bool:
    """Vrai si la colonne N° contient un nombre (ligne de données, pas de légende)."""
    n = _to_int(val)
    return n is not None and n > 0


def _sheet_header_map(ws) -> dict[str, int]:
    """Retourne {header_normalisé: index_colonne} pour la 1re ligne."""
    first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return {str(h or "").strip().lower(): i for i, h in enumerate(first_row)}


# ── Parsing par onglet ────────────────────────────────────────────────────────

def _parse_kaspersky(ws) -> list[dict]:
    hm = _sheet_header_map(ws)
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not _is_valid_num(row[hm.get("n°", 0)] if hm else row[0]):
            continue
        code = _to_str(row[hm["code d'activation"]] if "code d'activation" in hm else None)
        if not code:
            continue
        records.append(dict(
            numero          = _to_int(row[hm.get("n°", 0)]),
            produit         = _to_str(row[hm["produit"]]) or "Kaspersky Plus",
            code_activation = code,
            nb_machines     = _to_int(row[hm.get("nb machines")] if "nb machines" in hm else None),
            date_debut      = _to_date(row[hm["date de début"]] if "date de début" in hm else None),
            date_expiration = _to_date(row[hm["date d'expiration"]] if "date d'expiration" in hm else None),
        ))
    return records


def _parse_machines(ws) -> list[dict]:
    hm = _sheet_header_map(ws)
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not _is_valid_num(row[hm.get("n°", 0)]):
            continue
        machine = _to_str(row[hm["machine"]] if "machine" in hm else None)
        code    = _to_str(row[hm["code d'activation"]] if "code d'activation" in hm else None)
        if not machine or not code:
            continue
        records.append(dict(
            numero          = _to_int(row[hm.get("n°", 0)]),
            machine         = machine,
            code_activation = code,
            date_expiration = _to_date(row[hm["date d'expiration"]] if "date d'expiration" in hm else None),
            statut          = _to_str(row[hm["statut"]] if "statut" in hm else None),
        ))
    return records


def _parse_m365_comptes(ws) -> list[dict]:
    hm = _sheet_header_map(ws)
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not _is_valid_num(row[hm.get("n°", 0)]):
            continue
        compte = _to_str(row[hm["compte organisateur"]] if "compte organisateur" in hm else None)
        if not compte:
            continue
        records.append(dict(
            numero                 = _to_int(row[hm.get("n°", 0)]),
            produit                = _to_str(row[hm["produit"]] if "produit" in hm else None) or "Microsoft 365",
            compte_organisateur    = compte,
            email_organisateur     = _to_str(row[hm["email organisateur"]] if "email organisateur" in hm else None),
            nb_utilisateurs        = _to_int(row[hm["nb utilisateurs"]] if "nb utilisateurs" in hm else None),
            places_libres          = _to_int(row[hm["places libres"]] if "places libres" in hm else None),
            date_debut             = _to_date(row[hm["date de début"]] if "date de début" in hm else None),
            dernier_renouvellement = _to_date(row[hm["dernier renouvellement"]] if "dernier renouvellement" in hm else None),
            date_expiration        = _to_date(row[hm["date d'expiration"]] if "date d'expiration" in hm else None),
        ))
    return records


def _parse_m365_membres(ws) -> list[dict]:
    hm = _sheet_header_map(ws)
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        nom = _to_str(row[hm["nom"]] if "nom" in hm else None)
        if not nom:
            continue
        compte = _to_str(row[hm["compte organisateur"]] if "compte organisateur" in hm else None)
        if not compte:
            continue
        records.append(dict(
            numero              = _to_int(row[hm.get("n°", 0)]),
            nom                 = nom,
            email               = _to_str(row[hm["email"]] if "email" in hm else None),
            role                = _to_str(row[hm["rôle"]] if "rôle" in hm else (
                                  row[hm["role"]] if "role" in hm else None)),
            compte_organisateur = compte,
        ))
    return records


def _parse_adobe(ws) -> list[dict]:
    hm = _sheet_header_map(ws)
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not _is_valid_num(row[hm.get("n°", 0)]):
            continue
        produit = _to_str(row[hm["produit"]] if "produit" in hm else None)
        if not produit:
            continue
        records.append(dict(
            numero                 = _to_int(row[hm.get("n°", 0)]),
            produit                = produit,
            utilisateur            = _to_str(row[hm["utilisateur"]] if "utilisateur" in hm else None),
            email_compte           = _to_str(row[hm["email du compte adobe"]] if "email du compte adobe" in hm else None),
            fournisseur            = _to_str(row[hm["fournisseur"]] if "fournisseur" in hm else None),
            date_activation        = _to_date(row[hm["date d'activation"]] if "date d'activation" in hm else None),
            dernier_renouvellement = _to_date(row[hm["dernier renouvellement"]] if "dernier renouvellement" in hm else None),
            date_expiration        = _to_date(row[hm["date d'expiration"]] if "date d'expiration" in hm else None),
        ))
    return records


def _parse_powerbi(ws) -> list[dict]:
    hm = _sheet_header_map(ws)
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not _is_valid_num(row[hm.get("n°", 0)]):
            continue
        produit = _to_str(row[hm["produit"]] if "produit" in hm else None)
        if not produit:
            continue
        records.append(dict(
            numero                 = _to_int(row[hm.get("n°", 0)]),
            produit                = produit,
            utilisateur            = _to_str(row[hm["utilisateur"]] if "utilisateur" in hm else None),
            compte_powerbi         = _to_str(row[hm["compte power bi"]] if "compte power bi" in hm else None),
            fournisseur            = _to_str(row[hm["fournisseur"]] if "fournisseur" in hm else None),
            date_activation        = _to_date(row[hm["date d'activation"]] if "date d'activation" in hm else None),
            dernier_renouvellement = _to_date(row[hm["dernier renouvellement"]] if "dernier renouvellement" in hm else None),
            date_expiration        = _to_date(row[hm["date d'expiration"]] if "date d'expiration" in hm else None),
        ))
    return records


# ── Endpoint ──────────────────────────────────────────────────────────────────

@router.post("/import-fichier", status_code=200, summary="Importer le fichier Licences.xlsx (tous les onglets)")
async def import_licences_fichier(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: User = Depends(require_editor),
):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(400, "Seuls les fichiers .xlsx sont acceptés")

    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    sheets = {n.lower(): n for n in wb.sheetnames}

    stats: dict[str, int] = {}

    # ── 1. Kaspersky (licences + machines) ────────────────────────────────────
    ws_lic = wb[sheets["licences"]] if "licences" in sheets else None
    ws_mac = wb[sheets["machines"]] if "machines" in sheets else None

    kaspersky_records = _parse_kaspersky(ws_lic) if ws_lic else []
    machine_records   = _parse_machines(ws_mac)  if ws_mac  else []

    # Vider puis ré-insérer (les machines référencent les licences par FK)
    db.query(KasperskyMachine).delete()
    db.query(LicenceKaspersky).delete()
    db.flush()

    for rec in kaspersky_records:
        db.add(LicenceKaspersky(**rec))
    db.flush()

    # Vérifier que les codes des machines existent bien dans les licences
    valid_codes = {r["code_activation"] for r in kaspersky_records}
    for rec in machine_records:
        if rec["code_activation"] not in valid_codes:
            continue  # Code inconnu — ignorer plutôt que planter
        db.add(KasperskyMachine(**rec))

    stats["kaspersky_licences"] = len(kaspersky_records)
    stats["kaspersky_machines"]  = len(machine_records)

    # ── 2. Microsoft 365 (comptes + membres) ──────────────────────────────────
    ws_m365c = wb[sheets.get("m365 comptes", "")] if "m365 comptes" in sheets else None
    ws_m365m = wb[sheets.get("m365 membres", "")] if "m365 membres" in sheets else None

    m365_comptes  = _parse_m365_comptes(ws_m365c) if ws_m365c else []
    m365_membres  = _parse_m365_membres(ws_m365m) if ws_m365m else []

    db.query(LicenceM365Membre).delete()
    db.query(LicenceM365Compte).delete()
    db.flush()

    for rec in m365_comptes:
        db.add(LicenceM365Compte(**rec))
    db.flush()

    valid_comptes = {r["compte_organisateur"] for r in m365_comptes}
    for rec in m365_membres:
        if rec["compte_organisateur"] not in valid_comptes:
            continue
        db.add(LicenceM365Membre(**rec))

    stats["m365_comptes"] = len(m365_comptes)
    stats["m365_membres"] = len(m365_membres)

    # ── 3. Adobe ──────────────────────────────────────────────────────────────
    ws_adobe = wb[sheets["adobe"]] if "adobe" in sheets else None
    adobe_records = _parse_adobe(ws_adobe) if ws_adobe else []

    db.query(LicenceAdobe).delete()
    db.flush()
    for rec in adobe_records:
        db.add(LicenceAdobe(**rec))
    stats["adobe"] = len(adobe_records)

    # ── 4. Power BI ───────────────────────────────────────────────────────────
    ws_pbi = wb[sheets.get("power bi", "")] if "power bi" in sheets else None
    pbi_records = _parse_powerbi(ws_pbi) if ws_pbi else []

    db.query(LicencePowerBI).delete()
    db.flush()
    for rec in pbi_records:
        db.add(LicencePowerBI(**rec))
    stats["power_bi"] = len(pbi_records)

    db.commit()
    return {
        "message": "Import terminé",
        "stats":   stats,
        "total":   sum(stats.values()),
    }


# ── Endpoints de lecture ──────────────────────────────────────────────────────

@router.get("/kaspersky", summary="Licences Kaspersky avec leurs machines")
def list_kaspersky(db: Session = Depends(get_db), _: User = Depends(require_editor)):
    from sqlalchemy.orm import joinedload
    return (
        db.query(LicenceKaspersky)
        .options(joinedload(LicenceKaspersky.machines))
        .order_by(LicenceKaspersky.numero)
        .all()
    )


@router.get("/m365", summary="Comptes Microsoft 365 avec leurs membres")
def list_m365(db: Session = Depends(get_db), _: User = Depends(require_editor)):
    from sqlalchemy.orm import joinedload
    return (
        db.query(LicenceM365Compte)
        .options(joinedload(LicenceM365Compte.membres))
        .order_by(LicenceM365Compte.numero)
        .all()
    )


@router.get("/adobe", summary="Licences Adobe")
def list_adobe(db: Session = Depends(get_db), _: User = Depends(require_editor)):
    return db.query(LicenceAdobe).order_by(LicenceAdobe.numero).all()


@router.get("/powerbi", summary="Licences Power BI")
def list_powerbi(db: Session = Depends(get_db), _: User = Depends(require_editor)):
    return db.query(LicencePowerBI).order_by(LicencePowerBI.numero).all()
