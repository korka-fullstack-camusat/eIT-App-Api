"""
Import Excel du suivi du parc informatique.
Onglets traités :
  - Suivi Parc Informatique  (en-tête à la ligne 2 du fichier)
  - En Stock
"""
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from sqlalchemy.orm import Session
from typing import Optional
from datetime import datetime, date
import io, math

import openpyxl

from ..database import get_db
from ..models.suivi_parc import SuiviParcItem
from ..models.user import User
from ..services.auth_service import require_editor

router = APIRouter(prefix="/api/parc", tags=["Parc informatique (import)"])


# ── Utilitaires ───────────────────────────────────────────────────────────────

def _to_str(val) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    return None if s in ("", "nan", "None", "NaT") else s


def _to_date(val) -> date | None:
    if val is None:
        return None
    if isinstance(val, date) and not isinstance(val, datetime):
        return val
    if isinstance(val, datetime):
        return val.date()
    s = str(val).strip()
    if not s or s in ("nan", "None", "NaT"):
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    return None


def _to_po(val) -> str | None:
    """Le PO peut être un float (22292.0) → convertir en entier si possible."""
    if val is None:
        return None
    s = str(val).strip()
    if s in ("", "nan", "None", "NaT"):
        return None
    try:
        return str(int(float(s)))
    except (ValueError, TypeError):
        return s


def _sheet_col_map(headers: list) -> dict[str, int]:
    return {str(h or "").strip().lower(): i for i, h in enumerate(headers)}


# Colonnes Excel → champs du modèle
_COL_TO_FIELD = {
    "code ref":               "code_ref",
    "nom":                    "nom",
    "prenom":                 "prenom",
    "projet":                 "projet",
    "lieu":                   "lieu",
    "nature":                 "nature",
    "désignation equipement": "designation_equipement",
    "designation equipement": "designation_equipement",
    "ref carte réseau":       "ref_carte_reseau",
    "ref carte reseau":       "ref_carte_reseau",
    "n° serie":               "numero_serie",
    "n° série":               "numero_serie",
    "po":                     "po",
    "date d'attribution":     "date_attribution",
    "statut":                 "statut",
}

_DATE_FIELDS  = {"date_attribution"}
_PO_FIELDS    = {"po"}


def _parse_sheet(ws, sheet_name: str, data_row_start: int = 2) -> list[dict]:
    """
    Lit un onglet et retourne une liste de dicts prêts à insérer.
    data_row_start=2 → la ligne d'en-têtes est la 2e ligne (onglet principal).
    data_row_start=1 → la ligne d'en-têtes est la 1re ligne (onglet En Stock).
    """
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < data_row_start:
        return []

    header_row = rows[data_row_start - 1]
    col_map = _sheet_col_map(header_row)

    # Construire le mapping champ → index de colonne
    field_idx: dict[str, int] = {}
    for col_name, field in _COL_TO_FIELD.items():
        if col_name in col_map and field not in field_idx:
            field_idx[field] = col_map[col_name]

    records = []
    for row in rows[data_row_start:]:
        # Ignorer les lignes sans code_ref
        code_ref_idx = field_idx.get("code_ref")
        if code_ref_idx is None:
            continue
        code_ref = _to_str(row[code_ref_idx] if code_ref_idx < len(row) else None)
        if not code_ref:
            continue

        rec: dict = {"source_onglet": sheet_name}
        for field, idx in field_idx.items():
            raw = row[idx] if idx < len(row) else None
            if field in _DATE_FIELDS:
                rec[field] = _to_date(raw)
            elif field in _PO_FIELDS:
                rec[field] = _to_po(raw)
            else:
                rec[field] = _to_str(raw)

        records.append(rec)
    return records


# ── Endpoint d'import ─────────────────────────────────────────────────────────

@router.post(
    "/import",
    status_code=200,
    summary="Importer le fichier Suivi Parc Informatique (.xlsx)",
)
async def import_parc(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: User = Depends(require_editor),
):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(400, "Seuls les fichiers .xlsx sont acceptés")

    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    sheets_lower = {n.lower(): n for n in wb.sheetnames}

    all_records: list[dict] = []

    # Onglet principal : en-têtes en ligne 2
    main_key = next(
        (k for k in sheets_lower if "suivi" in k or "parc" in k),
        None,
    )
    if main_key:
        records = _parse_sheet(wb[sheets_lower[main_key]], sheets_lower[main_key], data_row_start=2)
        all_records.extend(records)

    # Onglet En Stock : en-têtes en ligne 1
    stock_key = next(
        (k for k in sheets_lower if "stock" in k),
        None,
    )
    if stock_key:
        records = _parse_sheet(wb[sheets_lower[stock_key]], sheets_lower[stock_key], data_row_start=1)
        all_records.extend(records)

    if not all_records:
        raise HTTPException(
            400,
            f"Aucune donnée trouvée. Onglets détectés : {wb.sheetnames}",
        )

    # Dédoublonner par code_ref (premier occurrence gagne)
    seen: set[str] = set()
    deduped = []
    for rec in all_records:
        cr = rec.get("code_ref", "")
        if cr and cr not in seen:
            seen.add(cr)
            deduped.append(rec)

    # Vider la table puis ré-insérer
    db.query(SuiviParcItem).delete()
    db.flush()
    for rec in deduped:
        db.add(SuiviParcItem(**rec))
    db.commit()

    return {
        "message": "Import terminé",
        "total_inserted": len(deduped),
        "suivi_parc": sum(1 for r in deduped if r.get("source_onglet") != "En Stock"),
        "en_stock": sum(1 for r in deduped if r.get("source_onglet") == "En Stock"),
    }


# ── Endpoints de lecture ──────────────────────────────────────────────────────

@router.get("/", summary="Liste du parc informatique importé")
def list_parc(
    nature:  Optional[str] = Query(None),
    statut:  Optional[str] = Query(None),
    lieu:    Optional[str] = Query(None),
    search:  Optional[str] = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(require_editor),
):
    q = db.query(SuiviParcItem)
    if nature:
        q = q.filter(SuiviParcItem.nature.ilike(f"%{nature}%"))
    if statut:
        q = q.filter(SuiviParcItem.statut.ilike(f"%{statut}%"))
    if lieu:
        q = q.filter(SuiviParcItem.lieu.ilike(f"%{lieu}%"))
    if search:
        from sqlalchemy import or_
        q = q.filter(or_(
            SuiviParcItem.nom.ilike(f"%{search}%"),
            SuiviParcItem.prenom.ilike(f"%{search}%"),
            SuiviParcItem.code_ref.ilike(f"%{search}%"),
            SuiviParcItem.designation_equipement.ilike(f"%{search}%"),
            SuiviParcItem.numero_serie.ilike(f"%{search}%"),
        ))
    return q.order_by(SuiviParcItem.code_ref).all()
